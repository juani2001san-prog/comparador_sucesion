# -*- coding: utf-8 -*-
"""
Conversor TXT → PS3 reemplazando codigos de cuenta por sus equivalencias
en el plan de cuentas destino.

Para cualquier cliente que exporte sus registraciones con el layout de 64
chars por linea: cada linea tiene la misma estructura en el archivo origen
y en el PS3 destino — lo unico que cambia son los chars 2-9 (codigo de
cuenta de 8 digitos).

Layout de la linea (64 chars):

    pos 1       : prefijo (siempre "1" en el ejemplo)
    pos 2-9     : codigo de cuenta origen (8 chars) — SE REEMPLAZA
    pos 10-17   : fecha YYYYMMDD
    pos 18-25   : numero de asiento
    pos 26-28   : renglon (3 chars)
    pos 29-64   : detalle + importe

Encoding: ISO-8859 (latin-1), terminador CRLF — se preserva.
"""

from __future__ import annotations

import io
from openpyxl import load_workbook


def leer_equivalencias(xlsx_bytes: bytes) -> tuple[dict[str, str], list[str]]:
    """
    Lee el Excel de equivalencias (hoja 1).

    Formato esperado (encabezados):
      CODIGO | CUENTA | CDOJWIN | ...

    Devuelve:
      (equiv, sin_destino)
      equiv       → {codigo_origen (8 chars): codigo_destino (8 chars)}
      sin_destino → lista de codigos que tienen fila pero no tienen CDOJWIN.
    """
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb.active

    # Localizar columnas CODIGO y CDOJWIN por nombre (case-insensitive).
    headers = [str(ws.cell(1, j).value or "").strip().upper()
               for j in range(1, ws.max_column + 1)]
    try:
        i_cod = headers.index("CODIGO") + 1
    except ValueError:
        i_cod = 1  # fallback: primera columna
    try:
        i_dst = headers.index("CDOJWIN") + 1
    except ValueError:
        i_dst = 3  # fallback: tercera columna

    equiv: dict[str, str] = {}
    sin_destino: list[str] = []

    for r in range(2, ws.max_row + 1):
        cod_raw = ws.cell(r, i_cod).value
        dst_raw = ws.cell(r, i_dst).value
        if cod_raw is None:
            continue
        cod = str(cod_raw).strip()
        if not cod:
            continue
        # Normalizo a 8 chars con padding de ceros a la izquierda (los
        # codigos vienen como numeros en el excel y perdieron el 0 inicial).
        cod = cod.zfill(8) if cod.isdigit() else cod[:8]
        dst = str(dst_raw or "").strip()
        if dst:
            dst = dst.zfill(8) if dst.isdigit() else dst[:8]
            equiv[cod] = dst
        else:
            sin_destino.append(cod)

    return equiv, sin_destino


def convertir(
    txt_bytes: bytes,
    equiv: dict[str, str],
    omitir_sin_equivalencia: bool = False,
) -> tuple[bytes, dict]:
    """
    Convierte el TXT reemplazando los chars 2-9 de cada linea por la
    equivalencia correspondiente del plan destino.

    Args:
        txt_bytes: contenido del TXT origen (ISO-8859).
        equiv: diccionario {codigo_origen: codigo_destino}, ambos 8 chars.
        omitir_sin_equivalencia:
            - False (default): si una cuenta no esta en el mapeo, la linea
              se incluye tal cual (sin reemplazar) y se registra en 'faltantes'.
            - True: la linea se omite del output y se registra en 'faltantes'.

    Returns:
        (ps3_bytes, stats)
        stats = {
          'lineas_entrada':  int,
          'lineas_salida':   int,
          'lineas_omitidas': int,
          'reemplazos':      int,
          'faltantes':       dict[str, int],  # codigo -> cuantas veces aparece
          'cortas':          int,              # lineas de menos de 9 chars (se preservan tal cual)
        }
    """
    texto = txt_bytes.decode("latin-1", errors="replace")

    salida: list[str] = []
    faltantes: dict[str, int] = {}
    lineas_in = 0
    lineas_omitidas = 0
    reemplazos = 0
    cortas = 0

    for linea in texto.splitlines():
        lineas_in += 1

        # Lineas de menos de 9 chars: no tienen codigo de cuenta — pasan tal cual.
        if len(linea) < 9:
            salida.append(linea)
            cortas += 1
            continue

        prefijo = linea[0]
        cuenta_orig = linea[1:9]
        resto = linea[9:]

        dst = equiv.get(cuenta_orig)
        if dst is None:
            faltantes[cuenta_orig] = faltantes.get(cuenta_orig, 0) + 1
            if omitir_sin_equivalencia:
                lineas_omitidas += 1
                continue
            # Si no se omite, dejamos la linea con la cuenta original.
            salida.append(linea)
        else:
            salida.append(prefijo + dst + resto)
            reemplazos += 1

    ps3 = "\r\n".join(salida)
    if salida:
        ps3 += "\r\n"

    return ps3.encode("latin-1", errors="replace"), {
        "lineas_entrada": lineas_in,
        "lineas_salida": len(salida),
        "lineas_omitidas": lineas_omitidas,
        "reemplazos": reemplazos,
        "faltantes": faltantes,
        "cortas": cortas,
    }
